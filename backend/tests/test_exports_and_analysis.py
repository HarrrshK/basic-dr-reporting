import csv
from datetime import date, timedelta
from io import BytesIO, StringIO

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook

from app.api.analytics import call_analysis
from app.api.doctors import create_doctor
from app.api.exports import calls_csv, calls_xlsx, doctors as export_doctors
from app.api.visits import create_visit
from app.models import Product
from app.schemas import DoctorCreate, VisitCreate


def collect_response(response):
    return response.body


def test_call_csv_and_excel_exports_include_call_details_and_apply_filters(db):
    today = date.today()
    doctor = create_doctor(DoctorCreate(name="Dr Call Export", area="KHOPOLI", hq="BADLAPUR",
        category="A", mobile="9876543210", existing_specialty="Cardiology",
        specialty_group="Cardio", clinic_hospital="Central Clinic"), db)
    other = create_doctor(DoctorCreate(name="Dr Other", area="KARJAT", hq="KARJAT", category="B"), db)
    product = Product(name="Export Product")
    db.add(product); db.commit()
    create_visit(VisitCreate(doctor_id=doctor.id, visit_date=today, visit_time="10:30",
        purpose="Product discussion", product_ids=[product.id], outcome="Interested", notes="Asked for follow-up",
        follow_up_required=True, follow_up_date=today + timedelta(days=2), follow_up_reason="Share study"), db)
    create_visit(VisitCreate(doctor_id=other.id, visit_date=today, purpose="Routine check"), db)

    response = calls_csv(date_from=today, date_to=today, area="KHOPOLI", purpose="Product discussion", db=db)
    rows = list(csv.reader(StringIO(collect_response(response).decode("utf-8-sig"))))
    assert len(rows) == 2
    assert rows[0][:6] == ["Visit ID", "Visit Date", "Visit Time", "Doctor ID", "Doctor Name", "Mobile"]
    assert rows[1][4] == "Dr Call Export"
    assert rows[1][13:18] == ["KHOPOLI", "BADLAPUR", "Product discussion", "Export Product", "Interested"]
    assert rows[1][18:21] == ["Asked for follow-up", "Yes", (today + timedelta(days=2)).isoformat()]

    excel = calls_xlsx(date_from=today, date_to=today, hq="BADLAPUR", product_id=product.id, db=db)
    workbook = load_workbook(BytesIO(excel.body), read_only=True)
    sheet = workbook["Call Reports"]
    assert sheet.max_row == 2
    assert sheet["E2"].value == "Dr Call Export"
    assert sheet["R2"].value == "Interested"


def test_doctor_export_search_uses_filtered_result_and_sanitizes_spreadsheet_formulas(db):
    create_doctor(DoctorCreate(name="=SUM(1,1)", mobile="12345", area="KHOPOLI"), db)
    create_doctor(DoctorCreate(name="Dr Unrelated", area="KARJAT"), db)
    response = export_doctors(q="12345", active=True, db=db)
    rows = list(csv.reader(StringIO(collect_response(response).decode("utf-8-sig"))))
    assert len(rows) == 2
    assert rows[1][2] == "'=SUM(1,1)"


def test_call_analysis_derives_coverage_frequency_gaps_and_products(db):
    today = date.today()
    frequent = create_doctor(DoctorCreate(name="Dr Frequent", area="KHOPOLI", hq="BADLAPUR", category="A"), db)
    never = create_doctor(DoctorCreate(name="Dr Never", area="KHOPOLI", hq="BADLAPUR", category="A"), db)
    inactive = create_doctor(DoctorCreate(name="Dr Inactive", area="KHOPOLI", hq="BADLAPUR", category="A", active=False), db)
    product = Product(name="Analysis Product"); db.add(product); db.commit()
    create_visit(VisitCreate(doctor_id=frequent.id, visit_date=today - timedelta(days=6), visit_time="10:00", purpose="Detailing",
        product_ids=[product.id], outcome="Interested"), db)
    create_visit(VisitCreate(doctor_id=frequent.id, visit_date=today, purpose="Detailing",
        product_ids=[product.id], outcome="Interested", follow_up_required=True,
        follow_up_date=today, follow_up_reason="Call back"), db)

    result = call_analysis(start=today - timedelta(days=6), end=today, area="KHOPOLI", hq="BADLAPUR",
                           category="A", doctor_id=None, db=db)
    assert result["summary"]["visits"] == 2
    assert result["summary"]["active_doctors"] == 2
    assert result["summary"]["active_doctors_visited"] == 1
    assert result["summary"]["coverage_percentage"] == 50
    assert result["frequency"]["once"] == 0
    assert result["frequency"]["two_to_five"] == 1
    frequent_result = next(item for item in result["doctors"] if item["id"] == frequent.id)
    assert frequent_result["average_gap_days"] == 6
    assert frequent_result["longest_gap_days"] == 6
    assert frequent_result["products_discussed"] == 1
    assert frequent_result["pending_follow_ups"] == 1
    assert result["follow_ups"]["due_today"] == 1
    assert result["products"][0]["visits"] == 2
    assert result["products"][0]["doctors"] == 1
    assert result["products"][0]["areas"] == 1
    assert result["time_distribution"][1] == {"name": "9 AM–12 PM", "visits": 1}
    assert result["purposes"] == [{"name": "Detailing", "visits": 2}]

    filtered = call_analysis(start=today - timedelta(days=6), end=today, area="KHOPOLI", hq="BADLAPUR",
        category="A", product_id=product.id, purpose="Detailing", follow_up_required=True, db=db)
    assert filtered["summary"]["visits"] == 1
    assert filtered["follow_ups"]["pending"] == 1


def test_call_analysis_rejects_reversed_date_range(db):
    with pytest.raises(HTTPException) as error:
        call_analysis(start=date.today(), end=date.today() - timedelta(days=1), db=db)
    assert error.value.status_code == 422
