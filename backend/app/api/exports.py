import csv
import json
from datetime import date
from io import BytesIO, StringIO

from fastapi import APIRouter, Depends
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import or_, select
from sqlalchemy.orm import Session, selectinload

from ..database import get_db
from ..models import Area, Doctor, Product, Visit

router = APIRouter(prefix="/exports", tags=["exports"])

DOCTOR_HEADERS = ["Doctor ID", "External ID", "Doctor Name", "Existing Specialty / Qualification",
                  "Specialty Group", "Qualification", "Area", "HQ", "Category", "Mobile",
                  "Gender", "Doctor Status", "Active", "Clinic/Hospital", "Additional Data"]
CALL_HEADERS = ["Visit ID", "Visit Date", "Visit Time", "Doctor ID", "Doctor Name", "Mobile",
                "Specialty", "Specialty Group", "Qualification", "Category", "Doctor Status",
                "Gender", "Clinic/Hospital", "Area", "HQ", "Purpose", "Products Discussed",
                "Outcome", "Visit Notes", "Follow-up Required", "Follow-up Date", "Follow-up Reason",
                "Follow-up Status"]


def safe_cell(value):
    if value is None:
        return ""
    if hasattr(value, "value"):
        value = value.value
    if isinstance(value, (date,)):
        return value.isoformat()
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, dict):
        value = json.dumps(value, ensure_ascii=False, sort_keys=True)
    text = str(value)
    # Stop spreadsheet programs interpreting imported notes/names as formulas.
    if text.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + text
    return text


def csv_response(name: str, headers: list[str], rows):
    output = StringIO(newline="")
    writer = csv.writer(output, dialect="excel")
    writer.writerow(headers)
    writer.writerows(([safe_cell(value) for value in row] for row in rows))
    content = output.getvalue().encode("utf-8-sig")
    return Response(content, media_type="text/csv; charset=utf-8",
                    headers={"Content-Disposition": f'attachment; filename="{name}"'})


def excel_response(name: str, headers: list[str], rows):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Call Reports"
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="176B50")
    row_count = 1
    widths = [len(header) for header in headers]
    for row in rows:
        cells = [safe_cell(value) for value in row]
        sheet.append(cells)
        row_count += 1
        for index, value in enumerate(cells):
            widths[index] = min(max(widths[index], len(value)), 52)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{sheet.cell(row=1, column=len(headers)).column_letter}{row_count}"
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[sheet.cell(row=1, column=index).column_letter].width = max(12, min(width + 2, 52))
    output = BytesIO()
    workbook.save(output)
    return Response(output.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="{name}"'})


@router.get("/doctors.csv")
def doctors(area: str | None = None, hq: str | None = None, category: str | None = None,
            specialty_group: str | None = None, active: bool | None = None, q: str | None = None,
            db: Session = Depends(get_db)):
    stmt = select(Doctor).options(selectinload(Doctor.area)).order_by(Doctor.name)
    if area:
        stmt = stmt.where(Doctor.area.has(Area.name == area))
    if hq:
        stmt = stmt.where(Doctor.hq == hq)
    if category:
        stmt = stmt.where(Doctor.category == category)
    if specialty_group:
        stmt = stmt.where(Doctor.specialty_group == specialty_group)
    if active is not None:
        stmt = stmt.where(Doctor.active == active)
    if q:
        term = f"%{q.strip()}%"
        stmt = stmt.where(or_(Doctor.name.ilike(term), Doctor.external_id.ilike(term),
                              Doctor.mobile.ilike(term), Doctor.clinic_hospital.ilike(term),
                              Doctor.existing_specialty.ilike(term), Doctor.area.has(Area.name.ilike(term)),
                              Doctor.hq.ilike(term)))
    records = db.scalars(stmt)
    rows = ((d.id, d.external_id, d.name, d.existing_specialty, d.specialty_group, d.qualification,
             d.area.name if d.area else "", d.hq, d.category, d.mobile, d.gender, d.doctor_status,
             d.active, d.clinic_hospital, d.extra_data or {}) for d in records)
    return csv_response("doctors.csv", DOCTOR_HEADERS, rows)


def call_records(db: Session, date_from: date | None = None, date_to: date | None = None,
                 area: str | None = None, hq: str | None = None, category: str | None = None,
                 doctor_id: int | None = None, product_id: int | None = None,
                 purpose: str | None = None, outcome: str | None = None,
                 follow_up_required: bool | None = None, q: str | None = None):
    stmt = (select(Visit).join(Doctor).options(
        selectinload(Visit.doctor).selectinload(Doctor.area), selectinload(Visit.products))
        .order_by(Visit.visit_date.desc(), Visit.visit_time.desc(), Visit.id.desc()))
    if date_from:
        stmt = stmt.where(Visit.visit_date >= date_from)
    if date_to:
        stmt = stmt.where(Visit.visit_date <= date_to)
    if area:
        stmt = stmt.where(Doctor.area.has(Area.name == area))
    if hq:
        stmt = stmt.where(Doctor.hq == hq)
    if category:
        stmt = stmt.where(Doctor.category == category)
    if doctor_id:
        stmt = stmt.where(Visit.doctor_id == doctor_id)
    if product_id:
        stmt = stmt.where(Visit.products.any(Product.id == product_id))
    if purpose:
        stmt = stmt.where(Visit.purpose == purpose)
    if outcome:
        stmt = stmt.where(Visit.outcome == outcome)
    if follow_up_required is not None:
        stmt = stmt.where(Visit.follow_up_required == follow_up_required)
    if q:
        term = f"%{q.strip()}%"
        stmt = stmt.where(or_(Doctor.name.ilike(term), Doctor.mobile.ilike(term),
                              Doctor.clinic_hospital.ilike(term), Doctor.existing_specialty.ilike(term),
                              Doctor.hq.ilike(term), Doctor.area.has(Area.name.ilike(term)),
                              Visit.purpose.ilike(term), Visit.outcome.ilike(term), Visit.notes.ilike(term),
                              Visit.products.any(Product.name.ilike(term))))
    return list(db.scalars(stmt).unique())


def call_rows(records):
    for visit in records:
        doctor = visit.doctor
        yield (visit.id, visit.visit_date, visit.visit_time, doctor.id, doctor.name, doctor.mobile,
               doctor.existing_specialty, doctor.specialty_group, doctor.qualification, doctor.category,
               doctor.doctor_status, doctor.gender, doctor.clinic_hospital,
               doctor.area.name if doctor.area else "", doctor.hq, visit.purpose,
               ", ".join(product.name for product in visit.products), visit.outcome, visit.notes,
               visit.follow_up_required, visit.follow_up_date, visit.follow_up_reason,
               visit.follow_up_status)


def export_calls(db: Session, kind: str, **filters):
    records = call_records(db, **filters)
    if kind == "xlsx":
        return excel_response("call-report.xlsx", CALL_HEADERS, call_rows(records))
    return csv_response("call-report.csv", CALL_HEADERS, call_rows(records))


@router.get("/calls.csv")
def calls_csv(date_from: date | None = None, date_to: date | None = None,
              area: str | None = None, hq: str | None = None, category: str | None = None,
              doctor_id: int | None = None, product_id: int | None = None,
              purpose: str | None = None, outcome: str | None = None,
              follow_up_required: bool | None = None, q: str | None = None,
              db: Session = Depends(get_db)):
    return export_calls(db, "csv", date_from=date_from, date_to=date_to, area=area, hq=hq,
                        category=category, doctor_id=doctor_id, product_id=product_id,
                        purpose=purpose, outcome=outcome,
                        follow_up_required=follow_up_required, q=q)


@router.get("/calls.xlsx")
def calls_xlsx(date_from: date | None = None, date_to: date | None = None,
               area: str | None = None, hq: str | None = None, category: str | None = None,
               doctor_id: int | None = None, product_id: int | None = None,
               purpose: str | None = None, outcome: str | None = None,
               follow_up_required: bool | None = None, q: str | None = None,
               db: Session = Depends(get_db)):
    return export_calls(db, "xlsx", date_from=date_from, date_to=date_to, area=area, hq=hq,
                        category=category, doctor_id=doctor_id, product_id=product_id,
                        purpose=purpose, outcome=outcome,
                        follow_up_required=follow_up_required, q=q)


@router.get("/visits.csv")
def visits_csv(date_from: date | None = None, date_to: date | None = None,
               area: str | None = None, hq: str | None = None, category: str | None = None,
               doctor_id: int | None = None, product_id: int | None = None,
               purpose: str | None = None, outcome: str | None = None,
               follow_up_required: bool | None = None, q: str | None = None,
               db: Session = Depends(get_db)):
    return export_calls(db, "csv", date_from=date_from, date_to=date_to, area=area, hq=hq,
                        category=category, doctor_id=doctor_id, product_id=product_id,
                        purpose=purpose, outcome=outcome,
                        follow_up_required=follow_up_required, q=q)
