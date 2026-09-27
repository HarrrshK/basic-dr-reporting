from __future__ import annotations

from hashlib import sha256
from io import BytesIO
from typing import Any
from uuid import uuid4

from openpyxl import load_workbook
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from ..models import Doctor, ImportBatch
from .doctors import apply_doctor_values, normalize_mobile, normalize_text


FIELD_ALIASES = {
    "external_id": {"doctor id", "dr id", "doctor code", "dr code"},
    "name": {"doctor name", "dr name", "doctor", "name"},
    "existing_specialty": {"existing specialty qualification", "existing specialty", "speciality", "specialty"},
    "area": {"area", "territory", "location"},
    "hq": {"hq", "headquarter", "headquarters"},
    "category": {"category", "class"},
    "mobile": {"mobile", "mobile number", "phone", "phone number", "contact"},
    "active": {"active", "is active", "active status"},
    "specialty_group": {"specialty group", "speciality group"},
    "doctor_status": {"doctor status", "status"},
    "qualification": {"qualification", "degree"},
    "gender": {"gender", "sex"},
    "clinic_hospital": {"clinic hospital name", "clinic hospital", "hospital name", "clinic name", "hospital", "clinic"},
}


def header_key(value: object) -> str:
    return " ".join("".join(c if c.isalnum() else " " for c in normalize_text(value)).split())


def auto_mapping(headers: list[str]) -> tuple[dict[str, str], list[str]]:
    mapping: dict[str, str] = {}
    uncertain: list[str] = []
    aliases = {alias: field for field, values in FIELD_ALIASES.items() for alias in values}
    for header in headers:
        key = header_key(header)
        if key in aliases:
            mapping[header] = aliases[key]
            continue
        uncertain.append(header)
    return mapping, uncertain


def parse_active(value: Any) -> bool | None:
    if value in (None, ""):
        return None
    key = normalize_text(value)
    if key in {"yes", "y", "true", "1", "active"}:
        return True
    if key in {"no", "n", "false", "0", "inactive"}:
        return False
    raise ValueError(f"Unrecognised active value: {value}")


def read_workbook(content: bytes) -> tuple[list[str], list[dict[str, Any]]]:
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    sheet = workbook.active
    iterator = sheet.iter_rows(values_only=True)
    first = next(iterator, None)
    if not first:
        raise ValueError("The workbook is empty")
    last_used = max((i for i, value in enumerate(first) if value not in (None, "")), default=-1)
    headers = [str(value).strip() if value is not None else f"Unnamed {i + 1}" for i, value in enumerate(first[:last_used + 1])]
    rows = []
    for excel_row, values in enumerate(iterator, start=2):
        record = {headers[i]: value for i, value in enumerate(values) if i < len(headers) and value not in (None, "")}
        if record:
            record["__row__"] = excel_row
            rows.append(record)
    return headers, rows


def mapped_values(raw: dict, mapping: dict[str, str]) -> tuple[dict, list[str]]:
    values: dict[str, Any] = {}
    errors: list[str] = []
    extra = {}
    for header, value in raw.items():
        if header == "__row__":
            continue
        field = mapping.get(header)
        if not field:
            extra[header] = value
        elif field == "active":
            try:
                values[field] = parse_active(value)
            except ValueError as exc:
                errors.append(str(exc))
        else:
            values[field] = str(int(value) if field == "mobile" and isinstance(value, float) and value.is_integer() else value).strip()
    values["extra_data"] = extra
    if not values.get("name"):
        errors.append("Doctor name is required")
    if values.get("external_id"):
        values["external_id"] = values["external_id"].upper()
    if "active" not in values and normalize_text(values.get("doctor_status")) in {"non active", "non-active", "inactive"}:
        values["active"] = False
    mobile = normalize_mobile(values.get("mobile"))
    if mobile and len(mobile) < 7:
        errors.append("Mobile number is too short")
    return values, errors


def is_id_only_template(values: dict) -> bool:
    return bool(values.get("external_id") and not values.get("name") and not values.get("extra_data") and
                all(value in (None, "") for key, value in values.items() if key not in {"external_id", "name", "extra_data"}))


def find_matches(db: Session, values: dict) -> list[Doctor]:
    external_id = normalize_text(values.get("external_id"))
    if external_id:
        exact = list(db.scalars(select(Doctor).where(func.lower(Doctor.external_id) == external_id)))
        if exact:
            return exact
    mobile = normalize_mobile(values.get("mobile"))
    if mobile:
        matches = list(db.scalars(select(Doctor).where(Doctor.normalized_mobile == mobile)))
        if external_id:
            # The supplied Doctor ID is authoritative. A shared mobile belonging
            # to another identified doctor must never merge these records.
            matches = [d for d in matches if not d.external_id or normalize_text(d.external_id) == external_id]
        if matches:
            return matches
    name = normalize_text(values.get("name"))
    if not name:
        return []
    candidates = list(db.scalars(select(Doctor).where(Doctor.normalized_name == name)))
    if external_id:
        candidates = [doctor for doctor in candidates if not doctor.external_id]
    clinic = normalize_text(values.get("clinic_hospital"))
    area = normalize_text(values.get("area"))
    specialty = normalize_text(values.get("existing_specialty"))
    strong = [d for d in candidates if sum([
        bool(clinic and normalize_text(d.clinic_hospital) == clinic),
        bool(area and d.area and normalize_text(d.area.name) == area),
        bool(specialty and normalize_text(d.existing_specialty) == specialty),
    ]) >= 2]
    return strong or candidates if len(candidates) > 1 else strong


def preview_import(db: Session, filename: str, content: bytes) -> ImportBatch:
    headers, rows = read_workbook(content)
    mapping, uncertain = auto_mapping(headers)
    preview_rows = []
    counts = {"new": 0, "matched": 0, "ambiguous": 0, "skipped": 0, "errors": 0}
    for raw in rows:
        values, errors = mapped_values(raw, mapping)
        placeholder = is_id_only_template(values)
        if placeholder:
            errors = []
        matches = find_matches(db, values) if not errors and not placeholder else []
        status = "skipped" if placeholder else "error" if errors else "new" if not matches else "matched" if len(matches) == 1 else "ambiguous"
        counts["errors" if status == "error" else status] += 1
        preview_rows.append({"row": raw["__row__"], "raw": raw, "values": values, "status": status,
                             "errors": errors, "matches": [{"id": d.id, "name": d.name} for d in matches]})
    batch = ImportBatch(id=str(uuid4()), filename=filename, file_hash=sha256(content).hexdigest(), mapping=mapping,
                        rows=preview_rows, summary={**counts, "total": len(rows), "unmapped_columns": uncertain})
    db.add(batch)
    db.commit()
    return batch


def confirm_import(db: Session, batch: ImportBatch, mapping: dict[str, str] | None, resolutions: dict[str, str]) -> dict:
    if batch.status != "preview":
        raise ValueError("This import has already been confirmed")
    used_mapping = mapping or batch.mapping
    summary = {"new": 0, "updated": 0, "ambiguous": 0, "skipped": 0, "errors": 0}
    for item in batch.rows:
        values, errors = mapped_values(item["raw"], used_mapping)
        if is_id_only_template(values):
            summary["skipped"] += 1
            continue
        if errors:
            summary["errors"] += 1
            continue
        matches = find_matches(db, values)
        resolution = resolutions.get(str(item["row"]))
        target = None
        if resolution == "create":
            target = None
        elif resolution and resolution.isdigit():
            target = db.get(Doctor, int(resolution))
        elif len(matches) == 1:
            target = matches[0]
        elif len(matches) > 1:
            summary["ambiguous"] += 1
            continue
        if target:
            apply_doctor_values(db, target, values, nonblank_only=True)
            summary["updated"] += 1
        else:
            doctor = Doctor(name=values["name"], normalized_name=normalize_text(values["name"]))
            apply_doctor_values(db, doctor, values)
            db.add(doctor)
            summary["new"] += 1
    batch.mapping = used_mapping
    batch.status = "completed"
    batch.summary = summary
    db.commit()
    return summary
